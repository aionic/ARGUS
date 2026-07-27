"""Render winner-only customer read-accuracy reports from cached ARGUS results."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from copy import copy
from datetime import UTC, datetime
from html import escape
from pathlib import Path
from statistics import mean
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table as ExcelTable
from openpyxl.worksheet.table import TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    TableStyle,
)
from reportlab.platypus import (
    Table as PdfTable,
)

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
GREEN = "C6EFCE"
RED = "FFC7CE"
YELLOW = "FFEB9C"
WHITE = "FFFFFF"

DISPLAY_NAMES = {
    "cms1500": "Structured (Claims)",
    "commercial-documents": "Invoices",
    "enrollments": "Handwriting/Handprint",
}


def _value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _mean(values: list[float]) -> float | None:
    return mean(values) if values else None


def _pct(value: Any) -> str:
    return "-" if not isinstance(value, (int, float)) else f"{value:.1%}"


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def _load_truth_metadata(corpus_root: Path) -> dict[str, dict[str, Any]]:
    metadata: dict[str, dict[str, Any]] = {}
    for dataset in ("commercial-documents", "enrollments"):
        path = corpus_root / "datasets" / dataset / "truth" / "records.jsonl"
        for record in _read_jsonl(path):
            metadata[record["document_id"]] = {
                "item_id": record.get("source_item_id"),
                "source_sheet": record.get("source_sheet"),
                "source_row": record.get("source_row"),
            }
    return metadata


def _load_field_labels(corpus_root: Path) -> dict[str, dict[str, str]]:
    schema_paths = {
        "cms1500": corpus_root / "schemas" / "cms1500.schema.json",
        "commercial-documents": corpus_root / "schemas" / "commercial-document.schema.json",
        "enrollments": corpus_root / "schemas" / "enrollment.schema.json",
    }
    labels: dict[str, dict[str, str]] = {}
    for dataset, path in schema_paths.items():
        schema = json.loads(path.read_text(encoding="utf-8"))
        labels[dataset] = {
            name: definition.get("x-source-label") or name for name, definition in schema.get("properties", {}).items()
        }
    return labels


def _variant_metrics(runs: list[dict[str, Any]]) -> tuple[float, float]:
    fields = [field for run in runs for field in (run.get("score") or {}).get("fields", [])]
    populated = [field for field in fields if field.get("expected_present")]
    blanks = [field for field in fields if field.get("expected_known") and not field.get("expected_present")]
    populated_accuracy = (
        sum(bool(field.get("normalized_match")) for field in populated) / len(populated) if populated else 0.0
    )
    blank_accuracy = sum(not field.get("actual_present") for field in blanks) / len(blanks) if blanks else 0.0
    return populated_accuracy, blank_accuracy


def _select_winners(report: dict[str, Any]) -> tuple[dict[str, str], list[dict[str, Any]]]:
    by_dataset_variant: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for run in report.get("runs", []):
        dataset = run.get("dataset")
        if dataset in DISPLAY_NAMES:
            by_dataset_variant[(dataset, run.get("variant", ""))].append(run)

    winners: dict[str, str] = {}
    selected: list[dict[str, Any]] = []
    for dataset in DISPLAY_NAMES:
        candidates = []
        for (candidate_dataset, variant), runs in by_dataset_variant.items():
            if candidate_dataset != dataset:
                continue
            populated_accuracy, blank_accuracy = _variant_metrics(runs)
            candidates.append((populated_accuracy, blank_accuracy, variant, runs))
        if not candidates:
            raise ValueError(f"No cached runs found for dataset '{dataset}'")
        _, _, winner, winner_runs = max(candidates, key=lambda candidate: candidate[:3])
        winners[dataset] = winner
        selected.extend(winner_runs)
    return winners, selected


def _document_identifier(run: dict[str, Any], truth_metadata: dict[str, dict[str, Any]]) -> str:
    metadata = truth_metadata.get(run.get("sample", ""), {})
    return metadata.get("item_id") or Path(run.get("original_filename", "")).stem


def _field_label(dataset: str, path: str, labels: dict[str, dict[str, str]]) -> str:
    root = path.split(".", 1)[0].split("[", 1)[0]
    label = labels.get(dataset, {}).get(root, root)
    suffix = path[len(root) :]
    return f"{label}{suffix}"


def _mean_field_confidence(run: dict[str, Any]) -> float | None:
    values = [
        float(field["confidence"])
        for field in (run.get("score") or {}).get("fields", [])
        if field.get("expected_present") and isinstance(field.get("confidence"), (int, float))
    ]
    return _mean(values)


def _grounding_coverage(run: dict[str, Any]) -> float | None:
    grounding = run.get("grounding") or {}
    extracted = int(grounding.get("extracted_fields") or 0)
    return int(grounding.get("grounded_fields") or 0) / extracted if extracted else None


def _field_status(field: dict[str, Any]) -> str:
    if field.get("normalized_match"):
        return "MATCH"
    if field.get("expected_present") and not field.get("actual_present"):
        return "MISSING"
    if not field.get("expected_present") and field.get("actual_present"):
        return "UNEXPECTED"
    return "MISMATCH"


def _style_sheet(sheet, freeze: str | None = None) -> None:
    sheet.sheet_view.showGridLines = False
    if freeze:
        sheet.freeze_panes = freeze
    for row in sheet.iter_rows():
        for cell in row:
            font = copy(cell.font)
            font.name = "Arial"
            cell.font = font
            cell.alignment = Alignment(vertical="top")


def _header(row) -> None:
    for cell in row:
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _add_table(sheet, name: str, end_column: str) -> None:
    table = ExcelTable(displayName=name, ref=f"A1:{end_column}{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _field_rows(
    runs: list[dict[str, Any]],
    truth_metadata: dict[str, dict[str, Any]],
    labels: dict[str, dict[str, str]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for run in runs:
        dataset = run["dataset"]
        for field in (run.get("score") or {}).get("fields", []):
            rows.append(
                {
                    "document_type": DISPLAY_NAMES[dataset],
                    "dataset": dataset,
                    "document_id": _document_identifier(run, truth_metadata),
                    "source_file": run.get("original_filename"),
                    "field_path": field.get("path"),
                    "field_label": _field_label(dataset, field.get("path", ""), labels),
                    "expected": field.get("expected"),
                    "actual": field.get("actual"),
                    "expected_present": bool(field.get("expected_present")),
                    "actual_present": bool(field.get("actual_present")),
                    "normalized_match": bool(field.get("normalized_match")),
                    "confidence": field.get("confidence"),
                }
            )
    return rows


def _build_workbook(
    report_title: str,
    runs: list[dict[str, Any]],
    truth_metadata: dict[str, dict[str, Any]],
    labels: dict[str, dict[str, str]],
    output: Path,
) -> None:
    field_rows = _field_rows(runs, truth_metadata, labels)
    workbook = Workbook()
    executive = workbook.active
    executive.title = "Executive Summary"
    executive["A1"] = report_title
    executive["A1"].font = Font(name="Arial", size=18, bold=True, color=WHITE)
    executive["A1"].fill = PatternFill("solid", fgColor=BLUE)
    executive.merge_cells("A1:F1")
    executive["A3"] = "Generated"
    executive["B3"] = datetime.now(UTC).isoformat()
    executive["A4"] = "Headline metric"
    executive["B4"] = "Normalized populated-field read accuracy"
    executive["A5"] = "Selection"
    executive["B5"] = "Highest all-document accuracy per document type; ties broken by blank accuracy"
    executive["A7"] = "Document Type"
    executive["B7"] = "Documents"
    executive["C7"] = "Correct Populated Fields"
    executive["D7"] = "Populated Golden Fields"
    executive["E7"] = "MS Read Rate Accuracy"
    executive["F7"] = "Source"
    _header(executive[7])

    document_sheet = workbook.create_sheet("Document Summary")
    document_columns = [
        "Document Type",
        "Document ID",
        "Source File",
        "Correct Populated Fields",
        "Populated Golden Fields",
        "Read Accuracy",
        "Correct Blank Fields",
        "Golden Blank Fields",
        "Blank Accuracy",
        "Mean Field Confidence",
        "OCR Word Confidence",
        "Grounding Coverage",
    ]
    document_sheet.append(document_columns)
    sorted_runs = sorted(
        runs,
        key=lambda run: (
            DISPLAY_NAMES[run["dataset"]],
            _document_identifier(run, truth_metadata),
        ),
    )
    for run in sorted_runs:
        metrics = (run.get("score") or {}).get("metrics") or {}
        row = document_sheet.max_row + 1
        document_sheet.append(
            [
                DISPLAY_NAMES[run["dataset"]],
                _document_identifier(run, truth_metadata),
                run.get("original_filename"),
                metrics.get("normalized_correct"),
                metrics.get("expected_fields"),
                f"=IFERROR(D{row}/E{row},0)",
                metrics.get("blank_correct"),
                metrics.get("expected_blank_fields"),
                f"=IFERROR(G{row}/H{row},0)",
                _mean_field_confidence(run),
                (run.get("ocr") or {}).get("mean"),
                _grounding_coverage(run),
            ]
        )
    _header(document_sheet[1])
    _style_sheet(document_sheet, "A2")
    for column in ("F", "I", "J", "K", "L"):
        for cell in document_sheet[column][1:]:
            cell.number_format = "0.0%"
    for column, width in zip(
        "ABCDEFGHIJKL",
        [25, 32, 36, 23, 23, 16, 20, 20, 16, 20, 20, 18],
        strict=True,
    ):
        document_sheet.column_dimensions[column].width = width
    _add_table(document_sheet, "DocumentAccuracyTable", "L")

    type_order = [
        name
        for name in ("Structured (Claims)", "Invoices", "Handwriting/Handprint")
        if any(DISPLAY_NAMES[run["dataset"]] == name for run in runs)
    ]
    for document_type in type_order:
        row = executive.max_row + 1
        executive.append(
            [
                document_type,
                f"=COUNTIF('Document Summary'!$A:$A,A{row})",
                f"=SUMIF('Document Summary'!$A:$A,A{row},'Document Summary'!$D:$D)",
                f"=SUMIF('Document Summary'!$A:$A,A{row},'Document Summary'!$E:$E)",
                f"=IFERROR(C{row}/D{row},0)",
                "Frozen ARGUS evaluation results",
            ]
        )
    overall_row = executive.max_row + 1
    first_type_row = 8
    last_type_row = overall_row - 1
    executive.append(
        [
            "Overall",
            f"=SUM(B{first_type_row}:B{last_type_row})",
            f"=SUM(C{first_type_row}:C{last_type_row})",
            f"=SUM(D{first_type_row}:D{last_type_row})",
            f"=IFERROR(C{overall_row}/D{overall_row},0)",
            "Weighted across populated golden fields",
        ]
    )
    for cell in executive[overall_row]:
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    for cell in executive["E"][7:]:
        cell.number_format = "0.0%"
    executive.column_dimensions["A"].width = 28
    executive.column_dimensions["B"].width = 14
    executive.column_dimensions["C"].width = 24
    executive.column_dimensions["D"].width = 24
    executive.column_dimensions["E"].width = 22
    executive.column_dimensions["F"].width = 40
    executive["B4"].alignment = executive["B5"].alignment = Alignment(wrap_text=True)
    _style_sheet(executive, "A8")

    field_accuracy = workbook.create_sheet("Field Accuracy")
    field_accuracy.append(
        [
            "Document Type",
            "Field",
            "Documents With Populated Truth",
            "Correct Populated Values",
            "Populated Golden Values",
            "Field Read Accuracy",
            "Mean Confidence",
        ]
    )
    field_keys = sorted({(row["document_type"], row["field_label"]) for row in field_rows if row["expected_present"]})
    evidence_last_row = len(field_rows) + 1
    for document_type, field_label in field_keys:
        row = field_accuracy.max_row + 1
        field_accuracy.append(
            [
                document_type,
                field_label,
                (
                    f"=COUNTIFS('Field Evidence'!$A$2:$A${evidence_last_row},A{row},"
                    f"'Field Evidence'!$F$2:$F${evidence_last_row},B{row},"
                    f"'Field Evidence'!$I$2:$I${evidence_last_row},TRUE)"
                ),
                (
                    f"=COUNTIFS('Field Evidence'!$A$2:$A${evidence_last_row},A{row},"
                    f"'Field Evidence'!$F$2:$F${evidence_last_row},B{row},"
                    f"'Field Evidence'!$I$2:$I${evidence_last_row},TRUE,"
                    f"'Field Evidence'!$K$2:$K${evidence_last_row},TRUE)"
                ),
                f"=C{row}",
                f"=IFERROR(D{row}/E{row},0)",
                (
                    f"=IFERROR(AVERAGEIFS('Field Evidence'!$M$2:$M${evidence_last_row},"
                    f"'Field Evidence'!$A$2:$A${evidence_last_row},A{row},"
                    f"'Field Evidence'!$F$2:$F${evidence_last_row},B{row},"
                    f"'Field Evidence'!$I$2:$I${evidence_last_row},TRUE),\"\")"
                ),
            ]
        )
    _header(field_accuracy[1])
    _style_sheet(field_accuracy, "A2")
    for column in ("F", "G"):
        for cell in field_accuracy[column][1:]:
            cell.number_format = "0.0%"
    for column, width in zip(
        "ABCDEFG",
        [25, 60, 28, 24, 24, 20, 18],
        strict=True,
    ):
        field_accuracy.column_dimensions[column].width = width
    _add_table(field_accuracy, "FieldAccuracyTable", "G")

    evidence = workbook.create_sheet("Field Evidence")
    evidence.append(
        [
            "Document Type",
            "Document ID",
            "Source File",
            "Field Path",
            "Field",
            "Field Label",
            "Golden Value",
            "Extracted Value",
            "Expected Present",
            "Actual Present",
            "Normalized Match",
            "Status",
            "Confidence",
            "Scored In Read Accuracy",
        ]
    )
    for source in field_rows:
        row = evidence.max_row + 1
        evidence.append(
            [
                source["document_type"],
                source["document_id"],
                source["source_file"],
                source["field_path"],
                source["field_path"],
                source["field_label"],
                _value(source["expected"]),
                _value(source["actual"]),
                source["expected_present"],
                source["actual_present"],
                source["normalized_match"],
                (
                    f'=IF(K{row},"MATCH",IF(AND(I{row}=TRUE,J{row}=FALSE),"MISSING",'
                    f'IF(AND(I{row}=FALSE,J{row}=TRUE),"UNEXPECTED",'
                    f'IF(AND(I{row}=FALSE,J{row}=FALSE),"BLANK OK","MISMATCH"))))'
                ),
                source["confidence"],
                source["expected_present"],
            ]
        )
    _header(evidence[1])
    _style_sheet(evidence, "A2")
    for column in ("D", "E", "F", "G", "H"):
        for cell in evidence[column][1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in evidence["M"][1:]:
        cell.number_format = "0.0%"
    widths = [25, 32, 36, 48, 48, 55, 42, 42, 17, 15, 18, 14, 14, 22]
    for index, width in enumerate(widths, start=1):
        evidence.column_dimensions[chr(64 + index)].width = width
    evidence.conditional_formatting.add(
        f"L2:L{evidence.max_row}",
        FormulaRule(formula=['L2="MATCH"'], fill=PatternFill("solid", fgColor=GREEN)),
    )
    evidence.conditional_formatting.add(
        f"L2:L{evidence.max_row}",
        FormulaRule(
            formula=['OR(L2="MISMATCH",L2="UNEXPECTED")'],
            fill=PatternFill("solid", fgColor=RED),
        ),
    )
    evidence.conditional_formatting.add(
        f"L2:L{evidence.max_row}",
        FormulaRule(formula=['L2="MISSING"'], fill=PatternFill("solid", fgColor=YELLOW)),
    )
    _add_table(evidence, "FieldEvidenceTable", "N")

    methodology = workbook.create_sheet("Methodology")
    methodology_rows = [
        ("Topic", "Definition"),
        (
            "MS Read Rate Accuracy",
            "Correct normalized populated-field values divided by all populated golden values.",
        ),
        (
            "Normalization",
            "Case, whitespace, punctuation, common numeric formatting, and null markers are normalized before comparison.",
        ),
        (
            "Winner selection",
            "Highest all-document populated-field micro accuracy per document type; blank accuracy breaks ties.",
        ),
        (
            "Blank fields",
            "Excluded from the headline read accuracy. Blank accuracy remains in supporting document detail.",
        ),
        (
            "Confidence",
            "Supporting Content Understanding field confidence and independent OCR word confidence; not part of accuracy.",
        ),
        (
            "Source",
            "Frozen cached ARGUS extraction results aligned to the canonical truth datasets.",
        ),
    ]
    for row in methodology_rows:
        methodology.append(row)
    _header(methodology[1])
    _style_sheet(methodology, "A2")
    methodology.column_dimensions["A"].width = 26
    methodology.column_dimensions["B"].width = 110
    for cell in methodology["B"]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def _pdf_text(value: Any, style: ParagraphStyle, limit: int = 180) -> Paragraph:
    text = escape(_value(value))
    if len(text) > limit:
        text = text[: limit - 3] + "..."
    return Paragraph(text or " ", style)


def _aggregate_type(runs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for run in runs:
        grouped[DISPLAY_NAMES[run["dataset"]]].append(run)
    rows = []
    for document_type in (
        "Structured (Claims)",
        "Invoices",
        "Handwriting/Handprint",
    ):
        type_runs = grouped.get(document_type, [])
        if not type_runs:
            continue
        fields = [field for run in type_runs for field in (run.get("score") or {}).get("fields", [])]
        populated = [field for field in fields if field.get("expected_present")]
        correct = sum(bool(field.get("normalized_match")) for field in populated)
        rows.append(
            {
                "document_type": document_type,
                "documents": len(type_runs),
                "correct": correct,
                "golden": len(populated),
                "accuracy": correct / len(populated) if populated else None,
            }
        )
    return rows


def _build_pdf(
    report_title: str,
    runs: list[dict[str, Any]],
    truth_metadata: dict[str, dict[str, Any]],
    labels: dict[str, dict[str, str]],
    output: Path,
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "Title",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor(f"#{BLUE}"),
        alignment=TA_CENTER,
        spaceAfter=14,
    )
    heading = ParagraphStyle(
        "Heading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=14,
        textColor=colors.HexColor(f"#{BLUE}"),
        alignment=TA_LEFT,
        spaceBefore=6,
        spaceAfter=6,
    )
    body = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
    )
    small = ParagraphStyle("Small", parent=body, fontSize=6.5, leading=8)
    document = SimpleDocTemplate(
        str(output),
        pagesize=landscape(letter),
        rightMargin=0.35 * inch,
        leftMargin=0.35 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.45 * inch,
        title=report_title,
        author="ARGUS",
    )
    story: list[Any] = [
        Paragraph(report_title, title),
        Paragraph(
            "Customer-facing read accuracy using only the best-performing cached output for each document type.",
            body,
        ),
        Spacer(1, 10),
    ]
    aggregate = _aggregate_type(runs)
    summary_data = [
        [
            "Document Type",
            "Documents",
            "Correct Populated Fields",
            "Populated Golden Fields",
            "MS Read Rate Accuracy",
        ]
    ]
    for row in aggregate:
        summary_data.append(
            [
                row["document_type"],
                row["documents"],
                row["correct"],
                row["golden"],
                _pct(row["accuracy"]),
            ]
        )
    total_correct = sum(row["correct"] for row in aggregate)
    total_golden = sum(row["golden"] for row in aggregate)
    summary_data.append(
        [
            "Overall",
            sum(row["documents"] for row in aggregate),
            total_correct,
            total_golden,
            _pct(total_correct / total_golden if total_golden else None),
        ]
    )
    summary_table = PdfTable(
        summary_data,
        colWidths=[2.2 * inch, 1.1 * inch, 1.7 * inch, 1.7 * inch, 1.6 * inch],
        repeatRows=1,
    )
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BLUE}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor(f"#{LIGHT_BLUE}")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story.extend(
        [
            summary_table,
            Spacer(1, 10),
            Paragraph(
                "Accuracy is calculated only on populated golden fields after normalization. Blank fields are excluded "
                "from the headline. Confidence is supporting information and does not affect correctness.",
                body,
            ),
            PageBreak(),
        ]
    )

    document_data = [
        [
            "Document Type",
            "Document ID",
            "Source File",
            "Correct",
            "Golden",
            "Read Accuracy",
            "Mean Confidence",
            "OCR Confidence",
        ]
    ]
    sorted_runs = sorted(
        runs,
        key=lambda run: (
            DISPLAY_NAMES[run["dataset"]],
            _document_identifier(run, truth_metadata),
        ),
    )
    for run in sorted_runs:
        metrics = (run.get("score") or {}).get("metrics") or {}
        document_data.append(
            [
                DISPLAY_NAMES[run["dataset"]],
                _document_identifier(run, truth_metadata),
                run.get("original_filename"),
                metrics.get("normalized_correct"),
                metrics.get("expected_fields"),
                _pct(metrics.get("normalized_accuracy")),
                _pct(_mean_field_confidence(run)),
                _pct((run.get("ocr") or {}).get("mean")),
            ]
        )
    document_table = PdfTable(
        document_data,
        colWidths=[
            1.6 * inch,
            1.8 * inch,
            2.4 * inch,
            0.65 * inch,
            0.65 * inch,
            0.9 * inch,
            0.95 * inch,
            0.9 * inch,
        ],
        repeatRows=1,
    )
    document_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BLUE}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 6.5),
                ("GRID", (0, 0), (-1, -1), 0.2, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor(f"#{LIGHT_BLUE}")],
                ),
            ]
        )
    )
    story.extend([Paragraph("Document-level summary", heading), document_table, PageBreak()])

    for index, run in enumerate(sorted_runs):
        dataset = run["dataset"]
        metrics = (run.get("score") or {}).get("metrics") or {}
        document_id = _document_identifier(run, truth_metadata)
        story.append(
            Paragraph(
                f"{escape(DISPLAY_NAMES[dataset])} - {escape(document_id)}",
                heading,
            )
        )
        story.append(
            Paragraph(
                f"Source: {escape(run.get('original_filename', ''))} | "
                f"Read accuracy: {_pct(metrics.get('normalized_accuracy'))} "
                f"({metrics.get('normalized_correct', 0)}/{metrics.get('expected_fields', 0)}) | "
                f"Mean field confidence: {_pct(_mean_field_confidence(run))}",
                body,
            )
        )
        mismatches = [
            field
            for field in (run.get("score") or {}).get("fields", [])
            if (field.get("expected_present") and not field.get("normalized_match"))
            or (not field.get("expected_present") and field.get("actual_present"))
        ]
        if not mismatches:
            story.append(Paragraph("No populated-field mismatches or unexpected extractions.", body))
        else:
            detail = [["Field", "Golden Value", "Extracted Value", "Confidence", "Result"]]
            for field in mismatches:
                detail.append(
                    [
                        _pdf_text(_field_label(dataset, field.get("path", ""), labels), small),
                        _pdf_text(field.get("expected"), small),
                        _pdf_text(field.get("actual"), small),
                        _pct(field.get("confidence")),
                        _field_status(field),
                    ]
                )
            detail_table = PdfTable(
                detail,
                colWidths=[
                    2.4 * inch,
                    2.55 * inch,
                    2.55 * inch,
                    0.8 * inch,
                    0.9 * inch,
                ],
                repeatRows=1,
            )
            detail_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BLUE}")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
                        ("GRID", (0, 0), (-1, -1), 0.2, colors.grey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        (
                            "ROWBACKGROUNDS",
                            (0, 1),
                            (-1, -1),
                            [colors.white, colors.HexColor(f"#{LIGHT_BLUE}")],
                        ),
                    ]
                )
            )
            story.append(detail_table)
        if index < len(sorted_runs) - 1:
            story.append(PageBreak())

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(0.4 * inch, 0.22 * inch, "ARGUS read accuracy - confidential local review")
        canvas.drawRightString(10.6 * inch, 0.22 * inch, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)


def _validate_workbook(path: Path, expected_documents: int) -> None:
    workbook = load_workbook(path, data_only=False, read_only=True)
    required = {
        "Executive Summary",
        "Document Summary",
        "Field Accuracy",
        "Field Evidence",
        "Methodology",
    }
    if set(workbook.sheetnames) != required:
        raise ValueError(f"Workbook sheets do not match expected output: {workbook.sheetnames}")
    if workbook["Document Summary"].max_row - 1 != expected_documents:
        raise ValueError("Workbook document count does not match the selected report scope")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("report_json", type=Path)
    parser.add_argument("--corpus-root", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    report = json.loads(args.report_json.read_text(encoding="utf-8"))
    winners, selected_runs = _select_winners(report)
    truth_metadata = _load_truth_metadata(args.corpus_root)
    labels = _load_field_labels(args.corpus_root)

    scopes = {
        "MS_Truth_Subset_Read_Accuracy": (
            "MS Truth Data - Invoice and Enrollment Read Accuracy",
            {"commercial-documents", "enrollments"},
        ),
        "ARGUS_All_Up_Read_Accuracy": (
            "ARGUS All-Up Read Accuracy - Best Output",
            {"cms1500", "commercial-documents", "enrollments"},
        ),
    }
    outputs = {}
    for filename, (title, datasets) in scopes.items():
        runs = [run for run in selected_runs if run["dataset"] in datasets]
        excel_path = args.output_dir / f"{filename}.xlsx"
        pdf_path = args.output_dir / f"{filename}.pdf"
        _build_workbook(title, runs, truth_metadata, labels, excel_path)
        _build_pdf(title, runs, truth_metadata, labels, pdf_path)
        _validate_workbook(excel_path, len(runs))
        outputs[filename] = {
            "excel": str(excel_path),
            "pdf": str(pdf_path),
            "documents": len(runs),
        }

    print(json.dumps({"winners": winners, "outputs": outputs}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
