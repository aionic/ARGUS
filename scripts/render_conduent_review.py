"""Render document-level Conduent bake-off results as Excel and PDF."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from datetime import UTC, datetime
from html import escape
from pathlib import Path
from statistics import mean
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table as ExcelTable
from openpyxl.worksheet.table import TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    TableStyle,
)
from reportlab.platypus import (
    Table as PdfTable,
)

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
GREEN = "C6EFCE"
RED = "FFC7CE"
YELLOW = "FFEB9C"
GRAY = "E7E6E6"
WHITE = "FFFFFF"


def _value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _pct(value: Any) -> str:
    return "-" if not isinstance(value, (int, float)) else f"{value:.1%}"


def _money(value: Any) -> str:
    return "-" if not isinstance(value, (int, float)) else f"${value:,.6f}"


def _field_status(field: dict[str, Any]) -> str:
    if field.get("normalized_match"):
        return "MATCH"
    if field.get("expected_present") and not field.get("actual_present"):
        return "MISSING"
    if not field.get("expected_present") and field.get("actual_present"):
        return "UNEXPECTED"
    return "MISMATCH"


def _mean_confidence(run: dict[str, Any]) -> float | None:
    values = [
        float(field["confidence"])
        for field in (run.get("score") or {}).get("fields", [])
        if isinstance(field.get("confidence"), (int, float))
    ]
    return mean(values) if values else None


def _style_sheet(sheet, freeze: str | None = None) -> None:
    sheet.sheet_view.showGridLines = False
    if freeze:
        sheet.freeze_panes = freeze
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="000000")
            cell.alignment = Alignment(vertical="top")


def _header(row) -> None:
    for cell in row:
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _table(sheet, name: str, start_row: int, end_row: int, end_column: str) -> None:
    if end_row < start_row:
        return
    table = ExcelTable(displayName=name, ref=f"A{start_row}:{end_column}{end_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _build_workbook(report: dict[str, Any], output: Path) -> None:
    workbook = Workbook()
    executive = workbook.active
    executive.title = "Executive Summary"
    totals = report.get("totals") or {}
    executive["A1"] = "ARGUS Conduent OCR and Extraction Bake-off"
    executive["A1"].font = Font(name="Arial", size=18, bold=True, color=WHITE)
    executive["A1"].fill = PatternFill("solid", fgColor=BLUE)
    executive.merge_cells("A1:H1")
    executive["A3"] = "Generated"
    executive["B3"] = datetime.now(UTC).isoformat()
    executive["A4"] = "Successful document-variant runs"
    executive["B4"] = totals.get("documents_succeeded")
    executive["A5"] = "Failed document-variant runs"
    executive["B5"] = totals.get("documents_failed")
    executive["A6"] = "Total measured processing cost"
    executive["B6"] = totals.get("total_cost_usd")
    executive["B6"].number_format = "$0.000000"
    executive["A8"] = "Pass gate"
    executive["B8"] = "At least 80% populated-field accuracy and 95% blank-field accuracy."
    executive["A10"] = "Contents"
    executive["B10"] = "Dataset Summary, Document Summary, Field Comparison, and Methodology"
    executive["A12"] = "Important"
    executive["B12"] = (
        "OCR confidence measures scan readability. Field confidence measures extraction certainty. "
        "Neither replaces comparison against the golden value."
    )
    executive.column_dimensions["A"].width = 34
    executive.column_dimensions["B"].width = 95
    executive["B8"].alignment = executive["B10"].alignment = executive["B12"].alignment = Alignment(wrap_text=True)
    _style_sheet(executive)

    dataset_sheet = workbook.create_sheet("Dataset Summary")
    dataset_columns = [
        "Dataset",
        "Split",
        "Variant",
        "Documents",
        "Passing",
        "Populated Accuracy",
        "Blank Accuracy",
        "Confidence Coverage",
        "OCR Confidence",
        "Grounding Coverage",
        "Total Cost (USD)",
        "Cost/Page (USD)",
        "P95 Latency (s)",
    ]
    dataset_sheet.append(dataset_columns)
    for summary in report.get("summaries", []):
        dataset_sheet.append(
            [
                summary.get("dataset"),
                summary.get("split"),
                summary.get("variant"),
                summary.get("documents_succeeded"),
                summary.get("documents_passing"),
                summary.get("normalized_accuracy"),
                summary.get("blank_accuracy"),
                summary.get("coverage"),
                summary.get("ocr_word_confidence_mean"),
                summary.get("grounding_coverage"),
                summary.get("total_cost_usd"),
                summary.get("cost_per_page_usd"),
                summary.get("p95_latency_seconds"),
            ]
        )
    _header(dataset_sheet[1])
    _style_sheet(dataset_sheet, "A2")
    dataset_sheet.auto_filter.ref = dataset_sheet.dimensions
    for column in ("F", "G", "H", "I", "J"):
        for cell in dataset_sheet[column][1:]:
            cell.number_format = "0.0%"
    for column in ("K", "L"):
        for cell in dataset_sheet[column][1:]:
            cell.number_format = "$0.000000"
    dataset_widths = [24, 14, 34, 12, 10, 18, 16, 20, 16, 18, 18, 18, 16]
    for index, width in enumerate(dataset_widths, start=1):
        dataset_sheet.column_dimensions[chr(64 + index)].width = width
    _table(dataset_sheet, "DatasetSummaryTable", 1, dataset_sheet.max_row, "M")

    field_sheet = workbook.create_sheet("Field Comparison")
    field_columns = [
        "Dataset",
        "Split",
        "Document ID",
        "Original File",
        "Variant",
        "Field Path",
        "Golden Value",
        "Extracted Value",
        "Expected Present",
        "Actual Present",
        "Exact Match",
        "Normalized Match",
        "Status",
        "Confidence",
    ]
    field_sheet.append(field_columns)
    runs = sorted(
        report.get("runs", []),
        key=lambda run: (
            run.get("dataset", ""),
            run.get("sample", ""),
            run.get("split", ""),
            run.get("variant", ""),
        ),
    )
    for run in runs:
        for field in (run.get("score") or {}).get("fields", []):
            row = field_sheet.max_row + 1
            field_sheet.append(
                [
                    run.get("dataset"),
                    run.get("split"),
                    run.get("sample"),
                    run.get("original_filename"),
                    run.get("variant"),
                    field.get("path"),
                    _value(field.get("expected")),
                    _value(field.get("actual")),
                    bool(field.get("expected_present")),
                    bool(field.get("actual_present")),
                    bool(field.get("exact_match")),
                    bool(field.get("normalized_match")),
                    (
                        f'=IF(L{row},"MATCH",IF(AND(I{row}=FALSE,J{row}=FALSE),"BLANK OK",'
                        f'IF(AND(I{row}=TRUE,J{row}=FALSE),"MISSING",'
                        f'IF(AND(I{row}=FALSE,J{row}=TRUE),"UNEXPECTED","MISMATCH"))))'
                    ),
                    field.get("confidence"),
                ]
            )
    _header(field_sheet[1])
    _style_sheet(field_sheet, "A2")
    field_sheet.auto_filter.ref = field_sheet.dimensions
    field_sheet.column_dimensions["A"].width = 24
    field_sheet.column_dimensions["B"].width = 14
    field_sheet.column_dimensions["C"].width = 28
    field_sheet.column_dimensions["D"].width = 34
    field_sheet.column_dimensions["E"].width = 34
    field_sheet.column_dimensions["F"].width = 48
    field_sheet.column_dimensions["G"].width = 42
    field_sheet.column_dimensions["H"].width = 42
    field_sheet.column_dimensions["I"].width = 16
    field_sheet.column_dimensions["J"].width = 14
    field_sheet.column_dimensions["K"].width = 13
    field_sheet.column_dimensions["L"].width = 18
    field_sheet.column_dimensions["M"].width = 14
    field_sheet.column_dimensions["N"].width = 13
    for row in field_sheet.iter_rows(min_row=2, min_col=6, max_col=8):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in field_sheet["N"][1:]:
        cell.number_format = "0.0%"
    field_sheet.conditional_formatting.add(
        f"M2:M{field_sheet.max_row}",
        FormulaRule(formula=['M2="MATCH"'], fill=PatternFill("solid", fgColor=GREEN)),
    )
    field_sheet.conditional_formatting.add(
        f"M2:M{field_sheet.max_row}",
        FormulaRule(formula=['OR(M2="MISMATCH",M2="UNEXPECTED")'], fill=PatternFill("solid", fgColor=RED)),
    )
    field_sheet.conditional_formatting.add(
        f"M2:M{field_sheet.max_row}",
        FormulaRule(formula=['M2="MISSING"'], fill=PatternFill("solid", fgColor=YELLOW)),
    )
    _table(field_sheet, "FieldComparisonTable", 1, field_sheet.max_row, "N")

    document_sheet = workbook.create_sheet("Document Summary")
    document_columns = [
        "Dataset",
        "Split",
        "Document ID",
        "Original File",
        "Variant",
        "Expected Populated Fields",
        "Normalized Correct",
        "Populated Accuracy",
        "Blank Accuracy",
        "Mean Field Confidence",
        "OCR Word Confidence",
        "Grounding Coverage",
        "Cost (USD)",
        "Latency (s)",
        "Gate",
    ]
    document_sheet.append(document_columns)
    last_field_row = field_sheet.max_row
    for run in runs:
        metrics = (run.get("score") or {}).get("metrics") or {}
        row = document_sheet.max_row + 1
        document_sheet.append(
            [
                run.get("dataset"),
                run.get("split"),
                run.get("sample"),
                run.get("original_filename"),
                run.get("variant"),
                metrics.get("expected_fields"),
                metrics.get("normalized_correct"),
                metrics.get("normalized_accuracy"),
                metrics.get("blank_accuracy"),
                (
                    f"=IFERROR(AVERAGEIFS('Field Comparison'!$N$2:$N${last_field_row},"
                    f"'Field Comparison'!$A$2:$A${last_field_row},A{row},"
                    f"'Field Comparison'!$B$2:$B${last_field_row},B{row},"
                    f"'Field Comparison'!$C$2:$C${last_field_row},C{row},"
                    f"'Field Comparison'!$E$2:$E${last_field_row},E{row}),\"\")"
                ),
                (run.get("ocr") or {}).get("mean"),
                (
                    (run.get("grounding") or {}).get("grounded_fields", 0)
                    / (run.get("grounding") or {}).get("extracted_fields", 1)
                    if (run.get("grounding") or {}).get("extracted_fields")
                    else None
                ),
                (run.get("cost") or {}).get("total_usd"),
                run.get("latency_seconds"),
                f'=IF(AND(H{row}>=0.8,OR(I{row}="",I{row}>=0.95)),"PASS","REVIEW")',
            ]
        )
    _header(document_sheet[1])
    _style_sheet(document_sheet, "A2")
    document_sheet.auto_filter.ref = document_sheet.dimensions
    for column in ("H", "I", "J", "K", "L"):
        for cell in document_sheet[column][1:]:
            cell.number_format = "0.0%"
    for cell in document_sheet["M"][1:]:
        cell.number_format = "$0.000000"
    document_widths = [24, 14, 28, 34, 34, 20, 18, 18, 16, 20, 18, 18, 15, 14, 12]
    for index, width in enumerate(document_widths, start=1):
        document_sheet.column_dimensions[chr(64 + index)].width = width
    document_sheet.conditional_formatting.add(
        f"O2:O{document_sheet.max_row}",
        CellIsRule(operator="equal", formula=['"PASS"'], fill=PatternFill("solid", fgColor=GREEN)),
    )
    document_sheet.conditional_formatting.add(
        f"O2:O{document_sheet.max_row}",
        CellIsRule(operator="equal", formula=['"REVIEW"'], fill=PatternFill("solid", fgColor=RED)),
    )
    _table(document_sheet, "DocumentSummaryTable", 1, document_sheet.max_row, "O")

    methodology = workbook.create_sheet("Methodology")
    methodology_rows = [
        ("Metric", "Definition"),
        ("Populated accuracy", "Normalized matches divided by populated golden fields."),
        ("Blank accuracy", "Golden blank fields that remained blank in extraction."),
        ("Field confidence", "Content Understanding confidence for the extracted field."),
        ("OCR confidence", "Independent Document Intelligence mean word confidence."),
        ("Grounding coverage", "Extracted fields with grounding metadata divided by extracted fields."),
        ("Pass gate", "Populated accuracy >= 80% and blank accuracy >= 95%."),
        ("Source", "Frozen ARGUS Conduent bake-off report.json and fields.csv."),
    ]
    for row in methodology_rows:
        methodology.append(row)
    _header(methodology[1])
    _style_sheet(methodology, "A2")
    methodology.column_dimensions["A"].width = 28
    methodology.column_dimensions["B"].width = 110
    for cell in methodology["B"]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def _pdf_paragraph(value: Any, style: ParagraphStyle) -> Paragraph:
    text = escape(_value(value))
    if len(text) > 180:
        text = text[:177] + "..."
    return Paragraph(text or " ", style)


def _build_pdf(report: dict[str, Any], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "ArgusTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor(f"#{BLUE}"),
        alignment=TA_CENTER,
        spaceAfter=14,
    )
    heading = ParagraphStyle(
        "ArgusHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=14,
        textColor=colors.HexColor(f"#{BLUE}"),
        alignment=TA_LEFT,
        spaceBefore=6,
        spaceAfter=6,
    )
    body = ParagraphStyle("ArgusBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=8, leading=10)
    small = ParagraphStyle("ArgusSmall", parent=body, fontSize=6.5, leading=8)
    document = SimpleDocTemplate(
        str(output),
        pagesize=landscape(letter),
        rightMargin=0.35 * inch,
        leftMargin=0.35 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.45 * inch,
        title="ARGUS Conduent Bake-off Document Review",
        author="ARGUS",
    )
    story: list[Any] = [
        Paragraph("ARGUS Conduent OCR and Extraction Bake-off", title),
        Paragraph(
            "Document-level review of extracted values, confidence, and comparison against the frozen golden datasets.",
            body,
        ),
        Spacer(1, 10),
    ]
    totals = report.get("totals") or {}
    overview = [
        ["Successful runs", "Failed runs", "Total measured cost", "Pass gate"],
        [
            totals.get("documents_succeeded", 0),
            totals.get("documents_failed", 0),
            _money(totals.get("total_cost_usd")),
            "80% populated / 95% blank",
        ],
    ]
    overview_table = PdfTable(
        overview,
        colWidths=[1.3 * inch, 1.1 * inch, 1.5 * inch, 2.2 * inch],
        repeatRows=1,
    )
    overview_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BLUE}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story.extend(
        [
            overview_table,
            Spacer(1, 10),
            Paragraph(
                "OCR confidence is an independent scan-readability measure. Field confidence is extraction certainty. "
                "The golden comparison is the correctness measure. Complete blank-field rows are retained in Excel but "
                "omitted from PDF detail when both golden and extracted values are blank.",
                body,
            ),
            PageBreak(),
        ]
    )

    grouped: dict[tuple[str, str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for run in report.get("runs", []):
        grouped[
            (
                run.get("dataset", ""),
                run.get("split", ""),
                run.get("sample", ""),
                run.get("original_filename", ""),
            )
        ].append(run)

    groups = sorted(grouped.items())
    for group_index, ((dataset, split, sample, original), runs) in enumerate(groups):
        story.append(Paragraph(f"{escape(dataset)} - {escape(sample)}", heading))
        story.append(Paragraph(f"Source file: {escape(original)} | Split: {escape(split)}", body))
        story.append(Spacer(1, 6))
        summary_data = [
            [
                "Variant",
                "Populated accuracy",
                "Blank accuracy",
                "Mean field confidence",
                "OCR confidence",
                "Grounding",
                "Cost",
                "Latency",
                "Gate",
            ]
        ]
        for run in sorted(runs, key=lambda item: item.get("variant", "")):
            metrics = (run.get("score") or {}).get("metrics") or {}
            grounding = run.get("grounding") or {}
            grounding_coverage = (
                grounding.get("grounded_fields", 0) / grounding.get("extracted_fields", 1)
                if grounding.get("extracted_fields")
                else None
            )
            passed = (
                isinstance(metrics.get("normalized_accuracy"), (int, float))
                and metrics["normalized_accuracy"] >= 0.80
                and (metrics.get("blank_accuracy") is None or metrics["blank_accuracy"] >= 0.95)
            )
            summary_data.append(
                [
                    run.get("variant"),
                    _pct(metrics.get("normalized_accuracy")),
                    _pct(metrics.get("blank_accuracy")),
                    _pct(_mean_confidence(run)),
                    _pct((run.get("ocr") or {}).get("mean")),
                    _pct(grounding_coverage),
                    _money((run.get("cost") or {}).get("total_usd")),
                    f"{float(run.get('latency_seconds') or 0):.2f}s",
                    "PASS" if passed else "REVIEW",
                ]
            )
        summary_table = PdfTable(
            summary_data,
            colWidths=[
                1.8 * inch,
                0.9 * inch,
                0.8 * inch,
                1.0 * inch,
                0.8 * inch,
                0.7 * inch,
                0.8 * inch,
                0.7 * inch,
                0.7 * inch,
            ],
            repeatRows=1,
        )
        summary_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BLUE}")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 6.5),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{LIGHT_BLUE}")]),
                ]
            )
        )
        story.extend([summary_table, Spacer(1, 8)])

        for run in sorted(runs, key=lambda item: item.get("variant", "")):
            variant = run.get("variant", "")
            story.append(Paragraph(f"Field comparison - {escape(variant)}", heading))
            detail_data: list[list[Any]] = [["Field", "Golden value", "Extracted value", "Confidence", "Result"]]
            fields = [
                field
                for field in (run.get("score") or {}).get("fields", [])
                if field.get("expected_present") or field.get("actual_present") or not field.get("normalized_match")
            ]
            for field in fields:
                detail_data.append(
                    [
                        _pdf_paragraph(field.get("path"), small),
                        _pdf_paragraph(field.get("expected"), small),
                        _pdf_paragraph(field.get("actual"), small),
                        _pct(field.get("confidence")),
                        _field_status(field),
                    ]
                )
            if len(detail_data) == 1:
                detail_data.append(["-", "-", "-", "-", "No populated fields"])
            detail_table = PdfTable(
                detail_data,
                colWidths=[2.35 * inch, 2.55 * inch, 2.55 * inch, 0.75 * inch, 0.9 * inch],
                repeatRows=1,
            )
            detail_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BLUE}")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
                        ("GRID", (0, 0), (-1, -1), 0.2, colors.grey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{LIGHT_BLUE}")]),
                    ]
                )
            )
            story.extend([detail_table, Spacer(1, 8)])
        if group_index < len(groups) - 1:
            story.append(PageBreak())

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(0.4 * inch, 0.22 * inch, "ARGUS Conduent bake-off - confidential local review")
        canvas.drawRightString(10.6 * inch, 0.22 * inch, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)


def _validate_workbook(path: Path) -> None:
    workbook = load_workbook(path, data_only=False, read_only=True)
    required = {"Executive Summary", "Dataset Summary", "Document Summary", "Field Comparison", "Methodology"}
    if set(workbook.sheetnames) != required:
        raise ValueError(f"Workbook sheets do not match expected output: {workbook.sheetnames}")
    if workbook["Document Summary"].max_row <= 1 or workbook["Field Comparison"].max_row <= 1:
        raise ValueError("Workbook is missing document or field detail rows")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("report_json", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    report = json.loads(args.report_json.read_text(encoding="utf-8"))
    excel_path = args.output_dir / "ARGUS_Conduent_Bakeoff_Document_Review.xlsx"
    pdf_path = args.output_dir / "ARGUS_Conduent_Bakeoff_Document_Review.pdf"
    _build_workbook(report, excel_path)
    _build_pdf(report, pdf_path)
    _validate_workbook(excel_path)
    print(json.dumps({"excel": str(excel_path), "pdf": str(pdf_path)}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
